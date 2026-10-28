import pandas as pd
import streamlit as st

from modules.auth import login, logout
from modules.config import DEMO_MODE
from modules.data import (
    CAREERS, CONTENT_STATUSES, QUESTION_TYPES, COMPETENCIES,
    get_careers, get_areas_for_career, get_courses_for_career,
    get_all_courses, get_all_areas, get_programs, get_topics_for_program,
    get_questions_for_career, get_course_careers, get_topics_for_course,
    get_all_questions, create_question, update_question,
    update_question_status, delete_question,
    build_simulator_question_set, create_attempt, get_attempt,
    get_attempt_answers, save_attempt_answer, finish_attempt,
    save_self_evaluation, recompute_attempt_score,
    get_student_attempts, get_attempt_diagnostics,
    EVALUATION_MODES, STRUCTURED_QUESTION_TYPES,
    get_all_questions_v16, create_question_v16, update_question_v16,
    get_quality_reviews, save_quality_review, bulk_transition_questions,
    get_published_questions_for_simulator_v16,
    build_simulator_question_set_v16, get_attempt_answers_v16,
    finalize_autograded_answers, save_self_evaluation_v16,
    recompute_attempt_score_v16, get_attempt_diagnostics_v16,
    review_answer_with_rubric, get_staff_attempts, get_profiles_by_ids,
    update_program_status, update_program_metadata,
    add_topic, update_topic, delete_topic,
)
from modules.ui import apply_branding, page_header, metric_card
from modules import grading_admin

st.set_page_config(
    page_title="PREP-EIME",
    page_icon="⚙️",
    layout="wide",
    initial_sidebar_state="expanded",
)
apply_branding()

STATUS_LABELS = {
    "borrador": "Borrador",
    "en_revision": "En revisión",
    "aprobado": "Aprobado",
    "publicado": "Publicado",
    "archivado": "Archivado",
}

CAREER_LABELS = {
    "IE": "Ingeniería Eléctrica",
    "IEL": "Ingeniería Electrónica",
    "IME": "Ingeniería Mecánica Eléctrica",
}


def login_view():
    c1, c2, c3 = st.columns([1, 1.3, 1])
    with c2:
        st.markdown("## PREP-EIME")
        st.caption("Plataforma de Repaso y Evaluación para el Examen General Privado")
        st.info("Acceso institucional restringido")
        email = st.text_input("Correo")
        password = st.text_input("Contraseña", type="password")
        if st.button("Ingresar", use_container_width=True, type="primary"):
            user = login(email, password)
            if user:
                st.session_state.user = user
                st.rerun()

        if DEMO_MODE:
            with st.expander("Credenciales de demostración"):
                st.code(
                    "estudiante@demo.gt / demo123\n"
                    "docente@demo.gt / demo123\n"
                    "coordinacion@demo.gt / demo123\n"
                    "admin@demo.gt / demo123"
                )


def sidebar(user):
    st.sidebar.markdown("## PREP-EIME")
    st.sidebar.caption("Escuela de Ingeniería Mecánica Eléctrica")
    st.sidebar.markdown("---")
    st.sidebar.write(f"**{user['name']}**")
    st.sidebar.caption(user["role"].capitalize())

    if user.get("career_code"):
        st.sidebar.caption(CAREERS.get(user["career_code"], user["career_code"]))
    if user.get("cohort"):
        st.sidebar.caption(user["cohort"])

    common = [
        "Inicio",
        "Preparación académica",
        "Biblioteca",
        "Preparación EGP",
        "Simulador EGP",
        "Mi progreso",
    ]

    if user["role"] == "docente":
        common += ["Gestión académica", "Banco de preguntas", "Control de calidad", "Revisión de respuestas"]
    elif user["role"] in ("coordinador", "administrador"):
        common += [
            "Gestión académica",
            "Banco de preguntas",
            "Control de calidad",
            "Revisión de respuestas",
            "Usuarios y cohortes",
            "Analítica",
        ]

    if user["role"] == "administrador":
        common.append("Evaluación administrativa")

    section = st.sidebar.radio(
        "Navegación", common, label_visibility="collapsed", key="main_navigation"
    )

    st.sidebar.markdown("---")
    st.sidebar.caption("PREP-EIME v1.7 · Control de calidad")
    if st.sidebar.button("Cerrar sesión"):
        logout()

    return section


def choose_career(user):
    if user.get("career_code"):
        return user["career_code"]

    careers = get_careers()
    mapping = {x["name"]: x["code"] for x in careers}
    if not mapping:
        return None

    name = st.selectbox("Carrera", list(mapping.keys()))
    return mapping[name]


def _flash(ok, msg):
    if ok:
        st.success(msg)
    else:
        st.error(msg)


def _allowed_transitions(role, status):
    if role == "docente":
        return {
            "borrador": ["en_revision"],
            "en_revision": ["borrador"],
        }.get(status, [])

    if role in ("coordinador", "administrador"):
        return {
            "borrador": ["en_revision"],
            "en_revision": ["borrador", "aprobado"],
            "aprobado": ["en_revision", "publicado"],
            "publicado": ["aprobado", "archivado"],
            "archivado": ["borrador"],
        }.get(status, [])

    return []


def home_view(user):
    page_header(f"Bienvenido, {user['name']}", f"Panel {user['role']}")

    career = choose_career(user) if user["role"] == "estudiante" else None

    if career:
        courses = get_courses_for_career(career)
        areas = get_areas_for_career(career)
        questions = get_questions_for_career(career)

        c1, c2, c3 = st.columns(3)
        with c1:
            metric_card("Cursos EGP", len(courses), "Catálogo institucional")
        with c2:
            metric_card("Áreas EGP", len(areas), "Según carrera")
        with c3:
            metric_card("Preguntas publicadas", len(questions), "Banco EGP")
    else:
        allc = get_all_courses()
        programs = get_programs(include_drafts=True)
        questions = get_all_questions(include_drafts=True)

        c1, c2, c3 = st.columns(3)
        with c1:
            metric_card("Cursos cargados", len(allc), "Supabase")
        with c2:
            metric_card("Programas", len(programs), "Incluye borradores")
        with c3:
            metric_card("Preguntas EGP", len(questions), "Banco institucional")

    st.success("PREP-EIME está conectado al catálogo académico real de Supabase.")


def academic_view(user):
    page_header("Preparación académica", "Ruta por carrera, área y curso")
    career = choose_career(user)

    if not career:
        st.warning("No hay carrera disponible para este perfil.")
        return

    courses = get_courses_for_career(career)
    if not courses:
        st.info("No se encontraron cursos EGP para esta carrera.")
        return

    df = pd.DataFrame(courses)[["code", "name", "semester", "area"]].rename(
        columns={
            "code": "Código",
            "name": "Curso",
            "semester": "Semestre",
            "area": "Área EGP",
        }
    )

    area = st.selectbox(
        "Filtrar por área",
        ["Todas"] + sorted(df["Área EGP"].dropna().unique().tolist()),
    )
    if area != "Todas":
        df = df[df["Área EGP"] == area]

    st.dataframe(df, use_container_width=True, hide_index=True)

    st.markdown("### Explorar contenido publicado")
    # Mantener el explorador sincronizado con el filtro de área.
    visible_codes = set(df["Código"].tolist())
    filtered_courses = [r for r in courses if r["code"] in visible_codes]
    opts = {f"{r['code']} · {r['name']}": r["code"] for r in filtered_courses}
    sel = st.selectbox("Curso", list(opts.keys()), key=f"academic_course_{career}_{area}")
    code = opts[sel]

    programs = [
        p for p in get_programs(include_drafts=False)
        if p["course_code"] == code
    ]

    if not programs:
        st.info("Este curso todavía no tiene un programa publicado para estudiantes.")
        return

    p = programs[0]
    st.caption(f"Programa: {p['version_label']} · Estado: {p['status']}")

    topics = get_topics_for_program(p["id"], include_drafts=False)
    if not topics:
        st.info("El programa está publicado, pero todavía no contiene unidades publicadas.")

    for t in topics:
        with st.expander(f"{t['sort_order']}. {t['title']}"):
            st.markdown(t.get("body_md") or "Sin desarrollo adicional.")


def library_view(user):
    page_header("Biblioteca académica", "Repositorio institucional")
    st.info(
        "El módulo de recursos se incorporará después de completar "
        "la revisión del banco de preguntas."
    )


def egp_view(user):
    page_header("Preparación EGP", "Preguntas publicadas para práctica")

    career = choose_career(user)
    questions = get_questions_for_career(career) if career else []

    if not questions:
        st.info("Todavía no hay preguntas publicadas para esta carrera.")
        return

    courses = {c["code"]: c for c in get_courses_for_career(career)}

    # Filtros dependientes: Área → Curso → Unidad temática → Dificultad.
    f1, f2, f3, f4 = st.columns(4)
    with f1:
        area_options = ["Todas"] + sorted({
            courses.get(q.get("course_code"), {}).get("area", "Sin área")
            for q in questions
        })
        area_filter = st.selectbox("Área", area_options, key=f"egp_area_{career}")

    area_questions = [
        q for q in questions
        if area_filter == "Todas"
        or courses.get(q.get("course_code"), {}).get("area", "Sin área") == area_filter
    ]
    with f2:
        course_options = ["Todos"] + sorted({
            q.get("course_code") for q in area_questions if q.get("course_code")
        })
        course_filter = st.selectbox(
            "Curso", course_options, key=f"egp_course_{career}_{area_filter}"
        )

    course_questions = [
        q for q in area_questions
        if course_filter == "Todos" or q.get("course_code") == course_filter
    ]

    # Mostrar solo unidades publicadas que realmente tengan preguntas publicadas
    # del curso seleccionado; no se mezclan unidades de distintos cursos.
    topic_labels = {}
    with f3:
        if course_filter != "Todos":
            linked_ids = {str(q["topic_id"]) for q in course_questions if q.get("topic_id")}
            for topic in get_topics_for_course(course_filter, include_drafts=False):
                topic_id = str(topic["id"])
                if topic_id in linked_ids:
                    label = f"{topic.get('sort_order', '—')}. {topic.get('title', 'Sin título')}"
                    # Identificador para distinguir nombres repetidos.
                    topic_labels[f"{label} · {topic_id[:8]}"] = topic_id
        topic_options = ["Todas"] + list(topic_labels)
        topic_filter = st.selectbox(
            "Unidad temática",
            topic_options,
            key=f"egp_topic_{career}_{area_filter}_{course_filter}",
            disabled=(course_filter == "Todos"),
            help="Selecciona primero un curso para consultar sus unidades."
        )

    with f4:
        difficulty_filter = st.selectbox(
            "Dificultad", ["Todas", 1, 2, 3, 4, 5],
            key=f"egp_difficulty_{career}"
        )

    selected_topic_id = topic_labels.get(topic_filter)
    filtered = [
        q for q in course_questions
        if (selected_topic_id is None or str(q.get("topic_id")) == selected_topic_id)
        and (difficulty_filter == "Todas" or q.get("difficulty") == difficulty_filter)
    ]

    st.caption(f"{len(filtered)} preguntas disponibles")

    for i, q in enumerate(filtered, 1):
        c = courses.get(q.get("course_code"), {})
        with st.expander(
            f"{i}. {q.get('course_code','')} · "
            f"{c.get('name','')} · Dificultad {q.get('difficulty','—')}/5"
        ):
            st.markdown(f"**Pregunta**")
            st.write(q["question_text"])
            st.caption(
                f"Tipo: {q.get('question_type','—')} · "
                f"Competencia: {q.get('competency','—')}"
            )

            show_key = f"show_answer_{q['id']}"
            if st.button("Mostrar respuesta esperada", key=show_key):
                st.info(q.get("expected_answer") or "No se registró una respuesta esperada.")



def _attempt_elapsed_minutes(attempt):
    from datetime import datetime, timezone
    if not attempt or not attempt.get("started_at"):
        return 0
    try:
        started = datetime.fromisoformat(attempt["started_at"].replace("Z", "+00:00"))
        return max(0, int((datetime.now(timezone.utc) - started).total_seconds() // 60))
    except Exception:
        return 0



def _mask_unpublished_administrative_scores(attempts):
    """Only the publication RPC authorizes student-visible administrative scores.

    This is UI-level defense in depth; database column permissions remain
    unchanged until staff-safe RPC reads are migrated.
    """
    for attempt in attempts:
        if attempt.get("grading_mode") != "administrativa":
            continue
        attempt["score"] = None  # Fail closed even if the RPC is unavailable.
        try:
            official = grading_admin.published_result(attempt["id"])
            if official and official.get("published"):
                attempt["score"] = official.get("score")
        except Exception:
            pass
    return attempts


def _attempt_history(user):
    """Muestra los intentos del estudiante sin alterar sus resultados."""
    if user.get("role") != "estudiante":
        return

    attempts = get_student_attempts(user.get("id")) or []
    st.markdown("### Historial de simulacros")
    if not attempts:
        st.info("Aún no tienes simulacros registrados.")
        return

    _mask_unpublished_administrative_scores(attempts)
    rows = []
    for attempt in attempts:
        raw_score = attempt.get("score")
        try:
            score = f"{float(raw_score):.1f}%" if raw_score is not None else "Pendiente"
        except (TypeError, ValueError):
            score = "Pendiente"
        rows.append({
            "Fecha": (attempt.get("started_at") or "")[:16].replace("T", " "),
            "Estado": attempt.get("status") or "—",
            "Preguntas": attempt.get("question_count") or 0,
            "Resultado": score,
        })
    st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)

    finished = [a for a in attempts if a.get("status") == "finalizado"]
    if finished:
        labels = {
            a["id"]: (
                f"{(a.get('started_at') or '')[:16].replace('T', ' ')} · "
                f"{a.get('question_count') or 0} preguntas · "
                f"{a.get('score') if a.get('score') is not None else 'Pendiente'}"
            )
            for a in finished
        }
        selected_id = st.selectbox(
            "Simulacro finalizado", list(labels),
            format_func=lambda attempt_id: labels[attempt_id],
            key="history_attempt_selection",
        )
        def open_selected_attempt():
            # Button callbacks execute before Streamlit instantiates the sidebar radio.
            st.session_state["review_attempt_id"] = selected_id
            st.session_state["main_navigation"] = "Simulador EGP"

        st.button(
            "Ver revisión del simulacro",
            key="open_attempt_history",
            on_click=open_selected_attempt,
        )


def simulator_view(user):
    page_header(
        "Simulador EGP",
        "Simulacro escrito con evaluación administrativa mediante rúbricas",
    )

    if user["role"] != "estudiante":
        st.info("El simulador está diseñado para perfiles de estudiante.")
        return

    career = user.get("career_code")
    if not career:
        st.warning("Tu perfil no tiene una carrera asignada.")
        return

    attempts = get_student_attempts(user["id"])
    active = next((a for a in attempts if a.get("status") == "en_progreso"), None)

    if active:
        try:
            active_rows = get_attempt_answers_v16(active["id"])
        except Exception as exc:
            st.error(f"No fue posible recuperar las preguntas: {exc}")
            return
        if not active_rows:
            st.error("Se encontró un intento en progreso sin preguntas recuperables.")
            st.caption(f"Intento: {active['id']}")
            if st.button("Cancelar intento incompleto y volver a configurar"):
                try:
                    from modules.supabase_client import get_supabase_client
                    client = get_supabase_client(require_auth=True)
                    response = (client.table("attempts").update({"status": "cancelado"})
                                .eq("id", active["id"]).eq("status", "en_progreso")
                                .execute())
                    if response.data:
                        st.session_state.pop("active_attempt_id", None)
                        st.session_state.pop("attempt_index", None)
                        st.rerun()
                    else:
                        st.error("No se pudo cancelar el intento.")
                except Exception as exc:
                    st.error(f"Error al cancelar: {exc}")
            return
        _run_active_attempt_v16(user, active)
        return

    published = get_published_questions_for_simulator_v16(career)
    if not published:
        st.warning("Todavía no hay preguntas publicadas para tu carrera.")
        return

    st.markdown("### Configurar nuevo simulacro")
    courses = {c["code"]: c for c in get_courses_for_career(career)}
    available_areas = sorted({
        courses.get(q.get("course_code"), {}).get("area", "Sin área")
        for q in published
    })

    c1, c2, c3 = st.columns(3)
    with c1:
        requested_count = st.number_input(
            "Número de preguntas", min_value=1,
            max_value=max(1, min(50, len(published))),
            value=min(10, len(published)), step=1
        )
    with c2:
        duration = st.number_input(
            "Tiempo sugerido (minutos)",
            min_value=5, max_value=240, value=60, step=5
        )
    with c3:
        difficulty_levels = st.multiselect(
            "Dificultad", [1,2,3,4,5], default=[1,2,3,4,5]
        )

    selected_areas = st.multiselect(
        "Áreas incluidas", available_areas, default=available_areas
    )
    st.caption("Las preguntas se distribuyen de forma equilibrada entre las áreas seleccionadas que tengan preguntas disponibles, según la dificultad elegida.")

    mode_counts = {}
    for q in published:
        mode = q.get("evaluation_mode") or "autoevaluacion"
        mode_counts[mode] = mode_counts.get(mode, 0) + 1

    st.caption(
        "Banco disponible: " +
        " · ".join(
            f"{EVALUATION_MODES.get(k,k)}: {v}"
            for k, v in sorted(mode_counts.items())
        )
    )

    # Validación previa: no registrar intentos sin consentimiento informado.
    # Se reutiliza exactamente el selector del simulador para respetar
    # carrera, áreas, dificultad y distribución equilibrada.
    if not selected_areas:
        st.warning("Selecciona al menos un área académica para continuar.")
        return
    if not difficulty_levels:
        st.warning("Selecciona al menos un nivel de dificultad para continuar.")
        return

    eligible_selection = build_simulator_question_set_v16(
        career, int(requested_count), selected_areas, difficulty_levels
    )
    available_count = len(eligible_selection)
    shortage = available_count < int(requested_count)

    if available_count == 0:
        st.error("No hay preguntas que cumplan los filtros seleccionados. Modifica las áreas o la dificultad.")
        return

    if shortage:
        st.warning(
            f"Solicitaste {int(requested_count)} preguntas, pero únicamente hay "
            f"{available_count} disponibles con los filtros seleccionados. "
            "Puedes ajustar los filtros o aceptar un simulacro reducido."
        )
        area_counts = {}
        for q in eligible_selection:
            area_name = q.get("area") or "Sin área"
            area_counts[area_name] = area_counts.get(area_name, 0) + 1
        st.caption("Disponibilidad por área: " + " · ".join(
            f"{area}: {area_counts.get(area, 0)}"
            for area in selected_areas
        ))
        start_label = f"Aceptar e iniciar con {available_count} preguntas"
    else:
        start_label = "Iniciar simulacro"

    if st.button(start_label, type="primary", use_container_width=True):
        ok, msg, attempt_id = create_attempt(
            user["id"], career,
            [q["id"] for q in eligible_selection], duration
        )
        _flash(ok, msg)
        if ok:
            st.session_state["active_attempt_id"] = attempt_id
            st.session_state["attempt_index"] = 0
            st.rerun()

    st.markdown("---")
    _attempt_history(user)


def _render_student_answer_widget(a):
    qtype = a.get("question_type")
    cfg = a.get("answer_config") or {}
    key = f"sim_answer_{a['id']}"
    current = st.session_state.get(key, a.get("answer_text") or "")

    if qtype == "Opción múltiple":
        options = cfg.get("options") or []
        display = ["— Selecciona una opción —"] + options
        initial = display.index(current) if current in display else 0
        selected = st.radio(
            "Tu respuesta", display, index=initial, key=key
        )
        return "" if selected.startswith("—") else selected

    if qtype == "Respuesta numérica":
        unit = cfg.get("unit") or ""
        value = st.text_input(
            f"Tu respuesta numérica{f' ({unit})' if unit else ''}",
            value=current,
            key=key,
            placeholder="Ejemplo: 12.75",
        )
        return value

    return st.text_area(
        "Tu respuesta",
        value=current,
        height=260,
        key=key,
        placeholder="Desarrolla tu procedimiento, análisis o justificación técnica.",
    )


def _run_active_attempt_v16(user, attempt):
    answers = get_attempt_answers_v16(attempt["id"])
    if not answers:
        st.error("El intento no contiene preguntas.")
        return

    idx = int(st.session_state.get("attempt_index", 0))
    idx = max(0, min(idx, len(answers)-1))
    st.session_state["attempt_index"] = idx

    elapsed = _attempt_elapsed_minutes(attempt)
    duration = attempt.get("duration_minutes") or 0
    remaining = max(0, duration - elapsed) if duration else None
    answered = sum(1 for a in answers if (a.get("answer_text") or "").strip())

    m1,m2,m3,m4 = st.columns(4)
    with m1:
        metric_card("Pregunta", f"{idx+1}/{len(answers)}", "Simulacro")
    with m2:
        metric_card("Respondidas", answered, f"de {len(answers)}")
    with m3:
        metric_card("Transcurrido", f"{elapsed} min", "Tiempo usado")
    with m4:
        metric_card("Restante", f"{remaining} min" if remaining is not None else "—", "Tiempo sugerido")

    a = answers[idx]
    st.markdown(f"### {a.get('course_code','')} · {a.get('course_name','')}")
    st.caption(
        f"Área: {a.get('area','Sin área')} · "
        f"Dificultad: {a.get('difficulty','—')}/5 · "
        f"Competencia: {a.get('competency','—')} · "
        f"Tipo: {a.get('question_type','—')} · "
        f"Evaluación: {EVALUATION_MODES.get(a.get('evaluation_mode'),'—')}"
    )
    st.markdown("#### Pregunta")
    st.write(a.get("question_text") or "")

    response = _render_student_answer_widget(a)

    if st.button("Guardar respuesta", use_container_width=True, key=f"save_answer_v16_{a['id']}"):
        ok, msg = save_attempt_answer(a["id"], response)
        _flash(ok, msg)
        if ok:
            st.rerun()

    nav1,nav2,nav3 = st.columns([1,1,2])
    with nav1:
        if st.button("← Anterior", disabled=idx==0, use_container_width=True):
            save_attempt_answer(a["id"], response)
            st.session_state["attempt_index"] = idx-1
            st.rerun()
    with nav2:
        if st.button("Siguiente →", disabled=idx==len(answers)-1, use_container_width=True):
            save_attempt_answer(a["id"], response)
            st.session_state["attempt_index"] = idx+1
            st.rerun()
    with nav3:
        confirm_finish = st.checkbox(
            "Confirmo que deseo entregar el simulacro",
            key=f"confirm_finish_v16_{attempt['id']}"
        )
        if st.button(
            "Finalizar y entregar",
            type="primary",
            disabled=not confirm_finish,
            use_container_width=True,
        ):
            save_attempt_answer(a["id"], response)
            ok, msg = finish_attempt(attempt["id"])
            if ok:
                if attempt.get("grading_mode") != "administrativa":
                    finalize_autograded_answers(attempt["id"])
            _flash(ok, msg)
            if ok:
                st.session_state.pop("attempt_index", None)
                st.session_state["review_attempt_id"] = attempt["id"]
                st.rerun()


def attempt_review_view(user):
    attempt_id = st.session_state.get("review_attempt_id")
    if not attempt_id:
        st.info("Selecciona un simulacro finalizado.")
        return

    attempt = get_attempt(attempt_id)
    answers = get_attempt_answers_v16(attempt_id)

    if attempt and attempt.get("grading_mode") == "administrativa":
        page_header("Revisión del simulacro", "Tus respuestas y soluciones esperadas")
        try:
            official = grading_admin.published_result(attempt_id)
        except Exception as exc:
            st.error(f"No fue posible consultar el estado de publicación: {exc}")
            official = None
        if official and official.get("published"):
            st.success("La administración publicó tu resultado oficial.")
            st.metric("Calificación oficial", f"{float(official['score']):.2f}/100")
        for i, a in enumerate(answers, 1):
            with st.expander(f"{i}. {a.get('course_name') or 'Pregunta'}"):
                st.markdown("**Pregunta**")
                st.write(a.get("question_text") or "")
                st.markdown("**Tu respuesta**")
                st.write(a.get("answer_text") or "Sin respuesta")
                st.markdown("**Solución esperada**")
                st.info(a.get("expected_answer") or "Sin solución registrada")
                if official and official.get("published"):
                    try:
                        details = grading_admin.published_rubric_details(attempt_id)
                        matching = next((d for d in details if d.get("answer_id") == a["id"]), None)
                        if matching:
                            st.markdown("**Evaluación por rúbrica**")
                            st.caption(f"Rúbrica: {matching.get('rubric_name', '—')}")
                            st.metric("Puntuación de esta respuesta",
                                f"{float(matching.get('answer_score') or 0):.2f}/100")
                            criteria_rows = matching.get("criteria") or []
                            if criteria_rows:
                                st.dataframe(pd.DataFrame([{
                                    "Criterio": r.get("criterion_name"),
                                    "Nivel": str(r.get("performance_level") or "").replace("_", " ").capitalize(),
                                    "Puntos": r.get("awarded_points"),
                                    "Máximo": r.get("max_points"),
                                } for r in criteria_rows]), use_container_width=True, hide_index=True)
                    except Exception:
                        st.caption("El desglose por criterios no está disponible temporalmente.")
        if not (official and official.get("published")):
            st.info("La calificación estará disponible cuando la administración la publique oficialmente.")
        if st.button("← Volver al simulador"):
            st.session_state.pop("review_attempt_id", None)
            st.rerun()
        return

    page_header(
        "Revisión del simulacro",
        "Resultados históricos: autoevaluación y revisión docente",
    )
    for i,a in enumerate(answers,1):
        mode = a.get("evaluation_mode") or "autoevaluacion"
        status = a.get("evaluation_status")

        with st.expander(
            f"{i}. {a.get('course_code','')} · {a.get('course_name','')} · "
            f"{EVALUATION_MODES.get(mode,mode)}"
        ):
            st.markdown("**Pregunta**")
            st.write(a.get("question_text") or "")
            st.markdown("**Tu respuesta**")
            st.write(a.get("answer_text") or "Sin respuesta.")

            if mode == "autocalificable":
                st.markdown("**Resultado automático**")
                if a.get("auto_score") is not None:
                    st.metric("Puntuación", f"{float(a['auto_score']):.0f}%")
                    st.caption(a.get("feedback") or "")
                st.markdown("**Explicación / respuesta esperada**")
                st.info(a.get("expected_answer") or "Sin explicación adicional.")

            elif mode == "autoevaluacion":
                st.markdown("**Respuesta esperada / criterio**")
                st.info(a.get("expected_answer") or "Sin criterio registrado.")
                score = st.slider(
                    "Porcentaje de cumplimiento",
                    0,100,
                    value=int(round(float(a.get("self_score") or 0))),
                    step=5,
                    key=f"self_score_v16_{a['id']}",
                )
                feedback = st.text_area(
                    "¿Qué debes reforzar?",
                    value=a.get("feedback") or "",
                    height=90,
                    key=f"self_feedback_v16_{a['id']}",
                )
                if st.button("Guardar autoevaluación", key=f"save_self_v16_{a['id']}"):
                    ok,msg = save_self_evaluation_v16(a["id"], score, feedback)
                    _flash(ok,msg)
                    if ok:
                        recompute_attempt_score_v16(attempt_id)
                        st.rerun()

            elif mode == "rubrica":
                st.markdown("**Estado**")
                if status == "revisado":
                    st.success(f"Revisada por el equipo académico · {float(a.get('review_score') or 0):.0f}%")
                    if a.get("feedback"):
                        st.write(a["feedback"])
                else:
                    st.warning("Pendiente de revisión docente.")
                if a.get("rubric"):
                    st.markdown("**Rúbrica de referencia**")
                    st.dataframe(pd.DataFrame(a["rubric"]), use_container_width=True, hide_index=True)

    # No recalcular notas durante una operación de lectura.
    diag = get_attempt_diagnostics_v16(attempt_id)

    st.markdown("---")
    st.markdown("## Estado del resultado")
    c1,c2,c3 = st.columns(3)
    with c1:
        metric_card(
            "Resultado actual",
            f"{diag['overall']:.1f}%" if diag["overall"] is not None else "—",
            "Solo respuestas evaluadas",
        )
    with c2:
        metric_card("Evaluadas", diag["evaluated"], f"de {diag['total']}")
    with c3:
        metric_card("Pendientes docente", diag["pending_review"], "Con rúbrica")

    if diag["evaluated"]:
        _diagnostic_panel_v16(attempt_id)

    if diag["pending_review"]:
        st.info(
            "El resultado puede cambiar cuando el docente revise las respuestas con rúbrica."
        )

    if st.button("← Volver al simulador"):
        st.session_state.pop("review_attempt_id", None)
        st.rerun()


def _diagnostic_panel_v16(attempt_id):
    diag = get_attempt_diagnostics_v16(attempt_id)
    if diag["overall"] is None:
        return

    st.markdown("### Diagnóstico")
    area_sorted = sorted(diag["by_area"], key=lambda x: x["score"], reverse=True)
    if len(area_sorted) == 1:
        area = area_sorted[0]
        metric_card("Desempeño del área", area["label"],
                    f"{area['score']:.1f}% · {area['n']} pregunta(s) evaluada(s)")
        st.caption("El diagnóstico se basa únicamente en las respuestas evaluadas; no representa necesariamente todas las preguntas del simulacro.")
    elif len(area_sorted) > 1:
        best, worst = area_sorted[0], area_sorted[-1]
        if abs(float(best["score"]) - float(worst["score"])) < 0.000001:
            st.info(f"Las {len(area_sorted)} áreas evaluadas presentan el mismo desempeño ({best['score']:.1f}%). No se identifica una fortaleza o área prioritaria diferenciada.")
        else:
            c1, c2 = st.columns(2)
            with c1:
                metric_card("Fortaleza", best["label"], f"{best['score']:.1f}%")
            with c2:
                metric_card("Área prioritaria", worst["label"], f"{worst['score']:.1f}%")
        st.caption("El diagnóstico se basa únicamente en las respuestas evaluadas.")

    if diag["by_area"]:
        df = pd.DataFrame(diag["by_area"]).rename(
            columns={"label":"Área","score":"Resultado (%)","n":"Preguntas"}
        )
        st.dataframe(df,use_container_width=True,hide_index=True)
        st.bar_chart(df.set_index("Área")["Resultado (%)"])

    if diag["by_competency"]:
        st.markdown("#### Por competencia")
        st.dataframe(
            pd.DataFrame(diag["by_competency"]).rename(
                columns={"label":"Competencia","score":"Resultado (%)","n":"Preguntas"}
            ),
            use_container_width=True, hide_index=True
        )

    if diag["by_difficulty"]:
        st.markdown("#### Por dificultad")
        st.dataframe(
            pd.DataFrame(diag["by_difficulty"]).rename(
                columns={"label":"Dificultad","score":"Resultado (%)","n":"Preguntas"}
            ),
            use_container_width=True, hide_index=True
        )


def progress_view(user):
    page_header("Mi progreso", "Historial y diagnóstico acumulado")

    if user["role"] != "estudiante":
        st.info("El progreso individual se muestra en perfiles de estudiante.")
        return

    attempts = _mask_unpublished_administrative_scores(
        get_student_attempts(user["id"])
    )
    completed = [a for a in attempts if a.get("status") == "finalizado"]

    if not completed:
        st.info("Aún no hay simulacros finalizados.")
        return

    scored = [a for a in completed if a.get("score") is not None]
    average = (
        sum(float(a["score"]) for a in scored) / len(scored)
        if scored else None
    )

    c1, c2, c3 = st.columns(3)
    with c1:
        metric_card("Simulacros", len(completed), "Finalizados")
    with c2:
        metric_card(
            "Promedio",
            f"{average:.1f}%" if average is not None else "—",
            "Intentos evaluados",
        )
    with c3:
        metric_card(
            "Preguntas practicadas",
            sum(int(a.get("question_count") or 0) for a in completed),
            "Acumulado",
        )

    rows = []
    for a in completed:
        rows.append({
            "Fecha": (a.get("started_at") or "")[:10],
            "Preguntas": a.get("question_count"),
            "Resultado": (
                f"{float(a['score']):.1f}%" if a.get("score") is not None
                else (
                    "Pendiente de publicación"
                    if a.get("grading_mode") == "administrativa"
                    and a.get("grading_status") in ("evaluado", "completo")
                    else "Pendiente de evaluación"
                    if a.get("grading_mode") == "administrativa"
                    else "Sin calificación"
                )
            ),
        })

    df = pd.DataFrame(rows)
    st.dataframe(df, use_container_width=True, hide_index=True)

    st.markdown("### Evolución del desempeño")
    evaluated_attempts = sorted(
        (a for a in completed if a.get("score") is not None),
        key=lambda a: a.get("started_at") or "",
    )
    if not evaluated_attempts:
        st.info("Aún no hay resultados evaluados para mostrar una evolución.")
    else:
        import plotly.graph_objects as go

        values = [float(a["score"]) for a in evaluated_attempts]
        labels = [f"Intento {i}" for i in range(1, len(values) + 1)]
        dates = [(a.get("started_at") or "")[:10] or "Sin fecha"
                 for a in evaluated_attempts]
        fig = go.Figure()
        fig.add_trace(go.Scatter(
            x=labels,
            y=values,
            mode="lines+markers" if len(values) > 1 else "markers",
            marker=dict(size=12),
            line=dict(width=3),
            customdata=dates,
            hovertemplate="%{x}<br>Fecha: %{customdata}<br>Resultado: %{y:.1f}%<extra></extra>",
            name="Resultado",
        ))
        fig.update_layout(
            height=320,
            margin=dict(l=10, r=10, t=15, b=10),
            showlegend=False,
            xaxis_title="Simulacro",
            yaxis_title="Resultado (%)",
            yaxis=dict(range=[0, 100], ticksuffix="%"),
        )
        st.plotly_chart(fig, use_container_width=True)
        if len(values) == 1:
            st.caption("Solo hay un simulacro evaluado. La tendencia se mostrará cuando completes otro.")

    st.markdown("---")
    _attempt_history(user)


def management_view(user):
    page_header("Gestión académica", "Edición, revisión, aprobación y publicación")

    staff = user["role"] in ("docente", "coordinador", "administrador")
    if not staff:
        st.warning("No tienes permisos de gestión académica.")
        return

    courses = get_all_courses()
    programs = get_programs(include_drafts=True)
    course_map = {c["code"]: c for c in courses}

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        metric_card("Cursos", len(courses), "Catálogo técnico")
    with c2:
        metric_card("Programas", len(programs), "Versiones registradas")
    with c3:
        metric_card(
            "Publicados",
            sum(1 for p in programs if p["status"] == "publicado"),
            "Visibles a estudiantes",
        )
    with c4:
        metric_card(
            "En revisión",
            sum(1 for p in programs if p["status"] == "en_revision"),
            "Pendientes de decisión",
        )

    tabs = st.tabs(["Mapa académico", "Editor de programa", "Editor de unidades"])

    with tabs[0]:
        if courses:
            rows = []
            p_by_code = {}
            for p in programs:
                p_by_code.setdefault(p["course_code"], []).append(p)

            for c in courses:
                ps = p_by_code.get(c["code"], [])
                rows.append({
                    "Código": c["code"],
                    "Curso": c["name"],
                    "Área EGP": c["area"],
                    "Semestre": c["semester"],
                    "Programas": len(ps),
                    "Estado": ", ".join(sorted({
                        STATUS_LABELS.get(p["status"], p["status"])
                        for p in ps
                    })) if ps else "Sin programa",
                })

            st.dataframe(
                pd.DataFrame(rows),
                use_container_width=True,
                hide_index=True,
            )

    with tabs[1]:
        if not programs:
            st.info("No hay programas disponibles.")
        else:
            labels = {}
            for p in programs:
                c = course_map.get(p["course_code"], {})
                labels[
                    f"{p['course_code']} · {c.get('name','')} · "
                    f"{STATUS_LABELS.get(p['status'],p['status'])}"
                ] = p

            label = st.selectbox(
                "Seleccionar programa",
                list(labels.keys()),
                key="editor_program",
            )
            p = labels[label]
            c = course_map.get(p["course_code"], {})

            st.markdown(f"### {p['course_code']} · {c.get('name','Curso')}")
            st.caption(
                f"Área: {c.get('area','Sin área')} · "
                f"Estado actual: {STATUS_LABELS.get(p['status'],p['status'])}"
            )

            with st.form("program_metadata"):
                version = st.text_input(
                    "Versión / identificación",
                    value=p.get("version_label") or "",
                )
                year = st.number_input(
                    "Año de vigencia",
                    min_value=1900,
                    max_value=2100,
                    value=int(p.get("effective_year") or 2026),
                    step=1,
                )
                source = st.text_area(
                    "Fuente / observación",
                    value=p.get("source_note") or "",
                    height=90,
                )
                save = st.form_submit_button(
                    "Guardar metadatos",
                    type="primary",
                )
                if save:
                    ok, msg = update_program_metadata(
                        p["id"], version, year, source
                    )
                    _flash(ok, msg)
                    if ok:
                        st.rerun()

            st.markdown("#### Flujo editorial")
            transitions = _allowed_transitions(user["role"], p["status"])

            if not transitions:
                st.caption(
                    "No hay transiciones disponibles para tu rol "
                    "desde el estado actual."
                )
            else:
                cols = st.columns(len(transitions))
                for col, target in zip(cols, transitions):
                    with col:
                        if st.button(
                            f"→ {STATUS_LABELS[target]}",
                            use_container_width=True,
                            key=f"pstatus_{p['id']}_{target}",
                        ):
                            ok, msg = update_program_status(p["id"], target)
                            _flash(ok, msg)
                            if ok:
                                st.rerun()

    with tabs[2]:
        if not programs:
            st.info("No hay programas disponibles.")
        else:
            labels = {}
            for p in programs:
                c = course_map.get(p["course_code"], {})
                labels[
                    f"{p['course_code']} · {c.get('name','')} · "
                    f"{STATUS_LABELS.get(p['status'],p['status'])}"
                ] = p

            label = st.selectbox(
                "Programa",
                list(labels.keys()),
                key="topic_program",
            )
            p = labels[label]
            topics = get_topics_for_program(p["id"], include_drafts=True)

            st.markdown("#### Unidades / temas existentes")
            if not topics:
                st.info("Este programa no tiene unidades registradas.")

            for t in topics:
                with st.expander(
                    f"{t['sort_order']}. {t['title']} · "
                    f"{STATUS_LABELS.get(t['status'],t['status'])}"
                ):
                    with st.form(f"topic_{t['id']}"):
                        title = st.text_input(
                            "Título",
                            value=t.get("title") or "",
                            key=f"title_{t['id']}",
                        )
                        order = st.number_input(
                            "Orden",
                            min_value=0,
                            max_value=999,
                            value=int(t.get("sort_order") or 0),
                            step=1,
                            key=f"order_{t['id']}",
                        )
                        body = st.text_area(
                            "Contenido",
                            value=t.get("body_md") or "",
                            height=260,
                            key=f"body_{t['id']}",
                        )

                        allowed_status = (
                            ["borrador", "en_revision"]
                            if user["role"] == "docente"
                            else list(STATUS_LABELS.keys())
                        )
                        current = t.get("status") or "borrador"
                        if current not in allowed_status:
                            allowed_status = [current] + [
                                x for x in allowed_status if x != current
                            ]

                        status = st.selectbox(
                            "Estado",
                            allowed_status,
                            index=allowed_status.index(current),
                            format_func=lambda x: STATUS_LABELS[x],
                            key=f"status_{t['id']}",
                        )

                        save_topic = st.form_submit_button(
                            "Guardar unidad / tema",
                            type="primary",
                        )

                        if save_topic:
                            ok, msg = update_topic(
                                t["id"], title, order, body, status
                            )
                            _flash(ok, msg)
                            if ok:
                                st.rerun()

                    if user["role"] == "administrador":
                        if st.button(
                            "Eliminar unidad / tema",
                            key=f"del_{t['id']}",
                        ):
                            ok, msg = delete_topic(t["id"])
                            _flash(ok, msg)
                            if ok:
                                st.rerun()

            st.markdown("#### Agregar nueva unidad / tema")
            with st.form("new_topic_form"):
                nt = st.text_input("Título de la nueva unidad")
                no = st.number_input(
                    "Orden",
                    min_value=0,
                    max_value=999,
                    value=(
                        max([
                            int(x.get("sort_order") or 0)
                            for x in topics
                        ], default=0) + 1
                    ),
                    step=1,
                )
                nb = st.text_area("Contenido inicial", height=180)
                submitted = st.form_submit_button(
                    "Crear unidad / tema",
                    type="primary",
                )

                if submitted:
                    ok, msg = add_topic(
                        p["id"], nt, no, nb, "borrador"
                    )
                    _flash(ok, msg)
                    if ok:
                        st.rerun()

    st.caption(
        "Regla editorial: docente prepara y envía a revisión; "
        "coordinación/administración aprueba y publica."
    )


def _build_answer_config(question_type, option_lines="", correct_option="", numeric_value=0.0, tolerance=0.0, unit=""):
    if question_type == "Opción múltiple":
        options = [x.strip() for x in (option_lines or "").splitlines() if x.strip()]
        return {"options": options, "correct_option": (correct_option or "").strip()}
    if question_type == "Respuesta numérica":
        return {
            "correct_value": float(numeric_value),
            "tolerance": abs(float(tolerance)),
            "unit": (unit or "").strip(),
        }
    return {}


def _parse_rubric_lines(text):
    rows = []
    for line in (text or "").splitlines():
        line = line.strip()
        if not line:
            continue
        if "|" in line:
            criterion, weight = line.rsplit("|", 1)
            try:
                w = float(weight.strip())
            except Exception:
                w = 1.0
            rows.append({"criterion": criterion.strip(), "weight": w})
        else:
            rows.append({"criterion": line, "weight": 1.0})
    return rows


def _rubric_to_text(rubric):
    return "\n".join(
        f"{x.get('criterion','')} | {x.get('weight',1)}"
        for x in (rubric or [])
    )


def question_bank_view(user):
    page_header(
        "Banco de preguntas EGP",
        "Preguntas autocalificables, abiertas y evaluadas con rúbrica",
    )

    if user["role"] not in ("docente", "coordinador", "administrador"):
        st.warning("No tienes permisos para administrar el banco.")
        return

    courses = get_all_courses()
    areas = get_all_areas()
    questions = get_all_questions_v16(include_drafts=True)
    course_map = {c["code"]: c for c in courses}
    area_map = {a["id"]: a["name"] for a in areas}

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        metric_card("Total", len(questions), "Preguntas registradas")
    with c2:
        metric_card(
            "Autocalificables",
            sum(1 for q in questions if q.get("evaluation_mode") == "autocalificable"),
            "Corrección automática",
        )
    with c3:
        metric_card(
            "Con rúbrica",
            sum(1 for q in questions if q.get("evaluation_mode") == "rubrica"),
            "Revisión académica",
        )
    with c4:
        metric_card(
            "Publicadas",
            sum(1 for q in questions if q["status"] == "publicado"),
            "Disponibles al estudiante",
        )

    tabs = st.tabs(["Explorar banco", "Crear pregunta", "Editar / flujo editorial"])

    with tabs[0]:
        if not questions:
            st.info("El banco todavía está vacío.")
        else:
            f1, f2, f3, f4 = st.columns(4)
            with f1:
                career_filter = st.selectbox(
                    "Carrera",
                    ["Todas", "IE", "IEL", "IME"],
                    format_func=lambda x: "Todas" if x == "Todas" else f"{x} · {CAREER_LABELS[x]}",
                    key="qb_filter_career",
                )
            with f2:
                area_filter = st.selectbox(
                    "Área EGP",
                    ["Todas"] + sorted(set(area_map.values())),
                    key="qb_filter_area",
                )
            with f3:
                mode_filter = st.selectbox(
                    "Evaluación",
                    ["Todas"] + list(EVALUATION_MODES.keys()),
                    format_func=lambda x: "Todas" if x == "Todas" else EVALUATION_MODES[x],
                    key="qb_filter_mode",
                )
            with f4:
                status_filter = st.selectbox(
                    "Estado",
                    ["Todos"] + list(STATUS_LABELS.keys()),
                    format_func=lambda x: "Todos" if x == "Todos" else STATUS_LABELS[x],
                    key="qb_filter_status",
                )

            rows = []
            filtered = []
            for q in questions:
                course = course_map.get(q.get("course_code"), {})
                qarea = area_map.get(q.get("area_id"), course.get("area", "Sin área"))
                careers = q.get("careers", [])

                if career_filter != "Todas" and career_filter not in careers:
                    continue
                if area_filter != "Todas" and qarea != area_filter:
                    continue
                if mode_filter != "Todas" and q.get("evaluation_mode") != mode_filter:
                    continue
                if status_filter != "Todos" and q["status"] != status_filter:
                    continue

                filtered.append(q)
                rows.append({
                    "Curso": q.get("course_code"),
                    "Carreras": ", ".join(careers),
                    "Tipo": q.get("question_type"),
                    "Evaluación": EVALUATION_MODES.get(q.get("evaluation_mode"), q.get("evaluation_mode")),
                    "Dificultad": q.get("difficulty"),
                    "Peso": q.get("weight"),
                    "Estado": STATUS_LABELS.get(q["status"], q["status"]),
                    "Pregunta": q["question_text"][:120],
                })

            st.caption(f"{len(filtered)} preguntas encontradas")
            if rows:
                st.dataframe(pd.DataFrame(rows), use_container_width=True, hide_index=True)

    with tabs[1]:
        if not courses:
            st.warning("No hay cursos disponibles.")
        else:
            course_labels = {f"{c['code']} · {c['name']}": c for c in courses}
            course_label = st.selectbox("Curso", list(course_labels.keys()), key="newq_course_v16")
            course = course_labels[course_label]
            linked_careers = get_course_careers(course["code"])
            topics = get_topics_for_course(course["code"], include_drafts=True)

            topic_labels = {"Sin tema específico": None}
            for t in topics:
                topic_labels[f"{t.get('sort_order',0)} · {t['title']}"] = t["id"]

            question_type = st.selectbox("Tipo de pregunta", STRUCTURED_QUESTION_TYPES, key="newq_type_v16")

            default_mode = (
                "autocalificable"
                if question_type in ("Opción múltiple", "Respuesta numérica")
                else ("rubrica" if question_type in ("Análisis de caso", "Diseño / selección", "Defensa oral") else "autoevaluacion")
            )
            allowed_modes = (
                ["autocalificable"]
                if question_type in ("Opción múltiple", "Respuesta numérica")
                else ["autoevaluacion", "rubrica"]
            )
            mode = st.selectbox(
                "Modalidad de evaluación",
                allowed_modes,
                index=allowed_modes.index(default_mode) if default_mode in allowed_modes else 0,
                format_func=lambda x: EVALUATION_MODES[x],
                key="newq_mode_v16",
            )

            with st.form("create_question_form_v16"):
                topic_label = st.selectbox("Unidad / tema", list(topic_labels.keys()))
                qtext = st.text_area("Enunciado de la pregunta", height=170)
                expected = st.text_area(
                    "Respuesta esperada / criterio general",
                    height=150,
                    help="Para preguntas autocalificables sirve como explicación posterior; para abiertas define el criterio técnico esperado.",
                )

                option_lines = ""
                correct_option = ""
                numeric_value = 0.0
                tolerance = 0.0
                unit = ""
                rubric_text = ""

                if question_type == "Opción múltiple":
                    option_lines = st.text_area(
                        "Opciones — una por línea",
                        height=120,
                        placeholder="Opción A\nOpción B\nOpción C\nOpción D",
                    )
                    correct_option = st.text_input(
                        "Respuesta correcta",
                        help="Debe coincidir exactamente con una de las opciones.",
                    )

                elif question_type == "Respuesta numérica":
                    n1, n2, n3 = st.columns(3)
                    with n1:
                        numeric_value = st.number_input("Valor correcto", value=0.0, format="%.8f")
                    with n2:
                        tolerance = st.number_input("Tolerancia absoluta ±", min_value=0.0, value=0.01, format="%.8f")
                    with n3:
                        unit = st.text_input("Unidad", placeholder="V, A, Ω, kW...")

                if mode == "rubrica":
                    rubric_text = st.text_area(
                        "Rúbrica",
                        height=140,
                        placeholder=(
                            "Identifica correctamente el principio físico | 30\n"
                            "Desarrolla el procedimiento técnico | 40\n"
                            "Justifica la conclusión | 30"
                        ),
                        help="Formato: criterio | peso. Los pesos son relativos.",
                    )

                r1, r2, r3, r4 = st.columns(4)
                with r1:
                    difficulty = st.select_slider("Dificultad", options=[1,2,3,4,5], value=3)
                with r2:
                    competency = st.selectbox("Competencia", COMPETENCIES)
                with r3:
                    weight = st.number_input("Peso", min_value=0.1, value=1.0, step=0.1)
                with r4:
                    careers = st.multiselect(
                        "Carreras",
                        ["IE","IEL","IME"],
                        default=linked_careers,
                        format_func=lambda x: f"{x} · {CAREER_LABELS[x]}",
                    )

                create = st.form_submit_button("Crear como borrador", type="primary", use_container_width=True)
                if create:
                    cfg = _build_answer_config(
                        question_type, option_lines, correct_option,
                        numeric_value, tolerance, unit
                    )

                    if question_type == "Opción múltiple":
                        options = cfg.get("options", [])
                        if len(options) < 2:
                            st.error("La opción múltiple necesita al menos dos opciones.")
                            st.stop()
                        if cfg.get("correct_option") not in options:
                            st.error("La respuesta correcta debe coincidir exactamente con una opción.")
                            st.stop()

                    rubric = _parse_rubric_lines(rubric_text) if mode == "rubrica" else []

                    ok, msg, _ = create_question_v16(
                        area_id=course.get("area_id"),
                        course_code=course["code"],
                        topic_id=topic_labels[topic_label],
                        question_text=qtext,
                        expected_answer=expected,
                        difficulty=difficulty,
                        competency=competency,
                        question_type=question_type,
                        career_codes=careers,
                        created_by=user["id"],
                        evaluation_mode=mode,
                        answer_config=cfg,
                        rubric=rubric,
                        weight=weight,
                    )
                    _flash(ok, msg)
                    if ok:
                        st.rerun()

    with tabs[2]:
        if not questions:
            st.info("No hay preguntas para editar.")
        else:
            labels = {}
            for q in questions:
                preview = q["question_text"].replace("\n", " ")[:75]
                labels[
                    f"{q.get('course_code','')} · "
                    f"{STATUS_LABELS.get(q['status'],q['status'])} · {preview}"
                ] = q

            selected_label = st.selectbox("Seleccionar pregunta", list(labels.keys()), key="edit_q_v16")
            q = labels[selected_label]
            course = course_map.get(q.get("course_code"), {})

            st.markdown(f"### {q.get('course_code','')} · {course.get('name','')}")
            st.caption(
                f"Evaluación: {EVALUATION_MODES.get(q.get('evaluation_mode'),'—')} · "
                f"Estado: {STATUS_LABELS.get(q['status'],q['status'])}"
            )

            st.markdown("**Enunciado actual**")
            st.write(q.get("question_text"))
            st.markdown("**Respuesta esperada**")
            st.write(q.get("expected_answer") or "—")

            cfg = q.get("answer_config") or {}
            if q.get("question_type") == "Opción múltiple":
                st.write("Opciones:", cfg.get("options", []))
                st.caption(f"Correcta: {cfg.get('correct_option','—')}")
            elif q.get("question_type") == "Respuesta numérica":
                st.caption(
                    f"Valor: {cfg.get('correct_value','—')} · "
                    f"Tolerancia: ±{cfg.get('tolerance','—')} · "
                    f"Unidad: {cfg.get('unit','')}"
                )

            if q.get("rubric"):
                st.markdown("**Rúbrica**")
                st.dataframe(pd.DataFrame(q["rubric"]), use_container_width=True, hide_index=True)

            st.markdown("#### Flujo editorial")
            transitions = _allowed_transitions(user["role"], q["status"])
            if transitions:
                cols = st.columns(len(transitions))
                for col, target in zip(cols, transitions):
                    with col:
                        if st.button(
                            f"→ {STATUS_LABELS[target]}",
                            use_container_width=True,
                            key=f"qstatus_v16_{q['id']}_{target}",
                        ):
                            ok, msg = update_question_status(q["id"], target)
                            _flash(ok, msg)
                            if ok:
                                st.rerun()

            st.info(
                "Para mantener la integridad de preguntas ya utilizadas en intentos, "
                "v1.6 prioriza el flujo editorial. Si necesitas cambiar sustancialmente "
                "una pregunta publicada, crea una nueva versión como borrador."
            )

            if user["role"] == "administrador":
                confirm = st.checkbox(
                    "Confirmo que deseo habilitar la eliminación",
                    key=f"confirm_delete_q_v16_{q['id']}",
                )
                if st.button(
                    "Eliminar pregunta",
                    disabled=not confirm,
                    key=f"delete_q_v16_{q['id']}",
                ):
                    ok, msg = delete_question(q["id"])
                    _flash(ok, msg)
                    if ok:
                        st.rerun()

    st.caption(
        "v1.6 distingue calificación automática, autoevaluación guiada "
        "y revisión académica mediante rúbrica."
    )



def _render_oral_evaluation_guide(config):
    """Presenta la configuración pedagógica sin exponer JSON técnico."""
    if not isinstance(config, dict) or not config:
        return

    st.markdown('#### Guía de evaluación oral')
    info = []
    if config.get('unit') is not None:
        info.append(f"Unidad {config['unit']}")
    if config.get('number') is not None:
        info.append(f"Pregunta {config['number']}")
    if config.get('modality'):
        info.append(f"Modalidad: {config['modality']}")
    if info:
        st.caption(' · '.join(info))

    if config.get('topic_block'):
        st.markdown('**Bloque temático**')
        st.write(str(config['topic_block']))
    if config.get('board'):
        st.markdown('**Desarrollo esperado en pizarrón**')
        st.write(str(config['board']))
    follow_up = config.get('follow_up')
    if isinstance(follow_up, (list, tuple)) and follow_up:
        st.markdown('**Preguntas complementarias**')
        for i, question in enumerate(follow_up, 1):
            st.write(f'{i}. {question}')
    elif isinstance(follow_up, str) and follow_up.strip():
        st.markdown('**Preguntas complementarias**')
        st.write(follow_up)
    if config.get('critical_error'):
        st.markdown('**Error crítico que debe identificarse**')
        st.warning(str(config['critical_error']))

    # Configuraciones de preguntas estructuradas, si se utilizan en este módulo.
    if isinstance(config.get('options'), list) and config['options']:
        st.markdown('**Opciones de respuesta**')
        for i, option in enumerate(config['options'], 1):
            st.write(f'{i}. {option}')
    for key, label in (
        ('correct_option', 'Opción correcta'),
        ('correct_value', 'Valor correcto'),
        ('tolerance', 'Tolerancia'),
        ('unit', 'Unidad de respuesta'),
    ):
        # unit numérico en bancos orales indica número de unidad temática.
        if key == 'unit' and isinstance(config.get(key), (int, float)):
            continue
        value = config.get(key)
        if value is not None and value != '':
            st.markdown(f'**{label}:** {value}')


def quality_management_view(user):
    page_header('Control de calidad EGP','Revisión académica y publicación masiva controlada')
    if user['role'] not in ('docente','coordinador','administrador'):
        st.error('Acceso restringido.'); return
    try:
        questions=get_all_questions_v16(include_drafts=True)
        reviews=get_quality_reviews()
        courses=get_all_courses()
    except Exception as exc:
        st.error(f'No se pudo leer la información. Ejecuta la migración v1.7: {exc}'); return
    course_names={c['code']:c['name'] for c in courses}
    m1,m2,m3=st.columns(3)
    m1.metric('Preguntas',len(questions))
    m2.metric('Revisadas',sum(bool(x.get('approved_for_publication')) for x in reviews.values()))
    m3.metric('Pendientes',sum(not reviews.get(q['id'],{}).get('approved_for_publication',False) for q in questions))
    c1,c2,c3=st.columns(3)
    with c1:
        course=st.selectbox('Curso',['Todos']+sorted({q['course_code'] for q in questions if q.get('course_code')}),format_func=lambda x:'Todos' if x=='Todos' else f'{x} · {course_names.get(x,x)}')
    with c2:
        status=st.selectbox('Estado',['Todos']+list(STATUS_LABELS),format_func=lambda x:'Todos' if x=='Todos' else STATUS_LABELS[x])
    with c3:
        quality=st.selectbox('Calidad',['Todas','Pendientes','Aprobadas'])
    filtered=[]
    for q in questions:
        approved=bool(reviews.get(q['id'],{}).get('approved_for_publication'))
        if course!='Todos' and q['course_code']!=course: continue
        if status!='Todos' and q['status']!=status: continue
        if quality=='Pendientes' and approved: continue
        if quality=='Aprobadas' and not approved: continue
        filtered.append(q)
    st.caption(f'{len(filtered)} preguntas encontradas')
    if not filtered: return
    t1,t2,t3=st.tabs(['Revisión individual','Transiciones masivas','Cobertura'])
    with t1:
        labels={f'{i+1}. {q.get("course_code")} · {q["question_text"][:80]}':q for i,q in enumerate(filtered)}
        selected=st.selectbox('Pregunta',list(labels)); q=labels[selected]
        st.markdown('**Pregunta**'); st.write(q['question_text'])
        st.markdown('**Respuesta esperada**'); st.write(q.get('expected_answer') or 'No registrada')
        st.caption(f'Unidad: {q.get("topic_id") or "No asignada"} · {q.get("question_type")} · Dificultad {q.get("difficulty")}')
        if q.get('rubric'): st.dataframe(pd.DataFrame(q['rubric']),hide_index=True)
        if q.get('answer_config'): _render_oral_evaluation_guide(q['answer_config'])
        old=reviews.get(q['id'],{})
        criteria=[('clarity','Enunciado claro'),('topic_alignment','Correspondencia con la unidad'),('technical_accuracy','Exactitud técnica'),('answer_quality','Respuesta correcta y suficiente'),('evaluation_quality','Rúbrica y modalidad apropiadas')]
        st.markdown('#### Evaluación de calidad académica')
        st.caption('Marque cada criterio únicamente después de verificarlo. Guardar la revisión no publica la pregunta.')
        with st.form('quality_review_form'):
            checks={k:st.checkbox(label,value=bool(old.get(k,False))) for k,label in criteria}
            fulfilled=sum(checks.values())
            st.progress(fulfilled / len(criteria), text=f'Criterios cumplidos: {fulfilled} de {len(criteria)}')
            if fulfilled == len(criteria):
                st.success('Los cinco criterios están marcados. La aprobación se determinará al guardar la revisión.')
            else:
                st.info('La pregunta conserva criterios pendientes de revisión.')
            notes=st.text_area('Observaciones académicas',value=old.get('notes') or '',placeholder='Describa errores detectados, ajustes necesarios o justificación académica...')
            if st.form_submit_button('Guardar revisión',type='primary'):
                ok,msg=save_quality_review(q['id'],user['id'],checks,notes)
                if ok: st.success(msg); st.rerun()
                else: st.error(msg)
    with t2:
        if user['role'] not in ('coordinador','administrador'):
            st.info('Las transiciones masivas requieren coordinación o administración.')
        else:
            approved=[q for q in filtered if reviews.get(q['id'],{}).get('approved_for_publication')]
            names={q['id']:f'{q["course_code"]} · {q["question_text"][:70]} · {q["status"]}' for q in approved}
            selected_ids=st.multiselect('Preguntas revisadas (máximo 100)',list(names),format_func=lambda x:names[x],max_selections=100)
            target=st.selectbox('Estado siguiente',['en_revision','aprobado','publicado'],format_func=lambda x:STATUS_LABELS[x])
            st.caption('Transiciones permitidas: borrador → en revisión → aprobado → publicado.')
            confirm=st.checkbox('Confirmo que revisé las preguntas seleccionadas')
            if st.button('Aplicar transición',type='primary',disabled=not selected_ids or not confirm):
                try:
                    result=bulk_transition_questions(selected_ids,target)
                    st.dataframe(pd.DataFrame(result),hide_index=True,use_container_width=True)
                    st.success(f'Actualizadas: {sum(r.get("outcome","").startswith("Actualizada") for r in result)}')
                except Exception as exc: st.error(str(exc))
    with t3:
        summary={}
        for q in questions:
            code=q.get('course_code') or 'Sin curso'
            item=summary.setdefault(code,{'Código':code,'Curso':course_names.get(code,code),'Total':0,'Revisadas':0,'Publicadas':0})
            item['Total']+=1
            item['Revisadas']+=bool(reviews.get(q['id'],{}).get('approved_for_publication'))
            item['Publicadas']+=q['status']=='publicado'
        df=pd.DataFrame(sorted(summary.values(),key=lambda x:x['Código']))
        st.dataframe(df,use_container_width=True,hide_index=True)
        st.download_button('Exportar CSV',df.to_csv(index=False).encode('utf-8-sig'),'cobertura_v1_7.csv','text/csv')

def review_answers_view(user):
    page_header(
        "Revisión de respuestas",
        "Calificación docente de respuestas abiertas y preguntas con rúbrica",
    )

    if user["role"] not in ("docente", "coordinador", "administrador"):
        st.warning("No tienes permisos para revisar respuestas.")
        return

    attempts = get_staff_attempts(limit=100)
    if not attempts:
        st.info("Todavía no hay simulacros finalizados.")
        return

    profiles = get_profiles_by_ids([a["student_id"] for a in attempts])

    pending_attempts = []
    for a in attempts:
        record = get_attempt(a["id"]) or {}
        if record.get("grading_mode") == "administrativa":
            continue
        answers = get_attempt_answers_v16(a["id"])
        pending = [
            x for x in answers
            if x.get("evaluation_mode") == "rubrica"
            and x.get("evaluation_status") != "revisado"
        ]
        if pending:
            pending_attempts.append((a, pending))

    c1, c2 = st.columns(2)
    with c1:
        metric_card("Intentos pendientes", len(pending_attempts), "Con respuestas por revisar")
    with c2:
        metric_card(
            "Respuestas pendientes",
            sum(len(x[1]) for x in pending_attempts),
            "Con rúbrica",
        )

    if not pending_attempts:
        st.success("No hay respuestas pendientes de revisión.")
        return

    labels = {}
    for a, pending in pending_attempts:
        p = profiles.get(a["student_id"], {})
        labels[
            f"{p.get('full_name','Estudiante')} · "
            f"{a.get('career_code','')} · "
            f"{(a.get('started_at') or '')[:10]} · "
            f"{len(pending)} pendiente(s)"
        ] = (a, pending)

    selected_label = st.selectbox("Seleccionar intento", list(labels.keys()))
    attempt, pending = labels[selected_label]
    profile = profiles.get(attempt["student_id"], {})

    st.markdown(f"### {profile.get('full_name','Estudiante')}")
    st.caption(
        f"Carné/ID: {profile.get('student_id') or '—'} · "
        f"Carrera: {attempt.get('career_code','—')}"
    )

    for i, a in enumerate(pending, 1):
        with st.expander(
            f"{i}. {a.get('course_code','')} · {a.get('course_name','')}",
            expanded=i == 1,
        ):
            st.markdown("**Pregunta**")
            st.write(a.get("question_text") or "")
            st.markdown("**Respuesta del estudiante**")
            st.write(a.get("answer_text") or "Sin respuesta.")

            st.markdown("**Respuesta esperada / criterio general**")
            st.info(a.get("expected_answer") or "Sin criterio general.")

            rubric = a.get("rubric") or []
            if rubric:
                st.markdown("**Rúbrica registrada**")
                st.dataframe(pd.DataFrame(rubric), use_container_width=True, hide_index=True)

            score = st.slider(
                "Calificación",
                0, 100,
                value=int(round(float(a.get("review_score") or 0))),
                step=5,
                key=f"review_score_{a['id']}",
            )
            feedback = st.text_area(
                "Retroalimentación docente",
                value=a.get("feedback") or "",
                height=120,
                key=f"review_feedback_{a['id']}",
            )

            if st.button("Guardar revisión", type="primary", key=f"review_save_{a['id']}"):
                ok, msg = review_answer_with_rubric(
                    a["id"], score, feedback, user["id"]
                )
                _flash(ok, msg)
                if ok:
                    recompute_attempt_score_v16(attempt["id"])
                    st.rerun()

def administrative_grading_view(user):
    page_header("Evaluación administrativa", "Rúbricas A, B y C · Evaluación por criterios")
    if user.get("role") != "administrador" or user.get("status") != "activo":
        st.error("Acceso exclusivo para administradores activos")
        return
    try:
        attempts = grading_admin.attempts()
        rubrics = grading_admin.rubrics()
        if not attempts or not rubrics:
            st.info("No hay simulacros administrativos finalizados o rúbricas activas.")
            return
        # Panel consolidado de evaluación y filtros de la bandeja administrativa.
        try:
            dashboard = grading_admin.dashboard_rows()
            st.subheader("Panel de gestión de evaluaciones")
            states = ("Pendiente", "En evaluación", "Evaluado", "Publicado")
            cols = st.columns(4)
            for col, state in zip(cols, states):
                col.metric(state, sum(r["estado_panel"] == state for r in dashboard))
            # Enriquecimiento de perfiles mediante la consulta ya existente, sin permisos nuevos.
            from datetime import datetime, timezone, timedelta, date
            from zoneinfo import ZoneInfo
            import io
            from collections import defaultdict
            guatemala_tz = ZoneInfo("America/Guatemala")
            profile_error = None
            try:
                student_profiles = get_profiles_by_ids([r["student_id"] for r in dashboard if r.get("student_id")])
            except Exception as exc:
                student_profiles = {}
                profile_error = str(exc)
            for r in dashboard:
                profile = student_profiles.get(r.get("student_id"), {})
                r["nombre_estudiante"] = profile.get("full_name") or "Sin nombre disponible"
                r["carne_estudiante"] = str(profile.get("student_id") or "")
                try:
                    dt = datetime.fromisoformat(str(r.get("finished_at") or "").replace("Z", "+00:00"))
                    if dt.tzinfo is None:
                        dt = dt.replace(tzinfo=timezone.utc)
                    r["fecha_local"] = dt.astimezone(guatemala_tz)
                except (ValueError, TypeError):
                    r["fecha_local"] = None
            if profile_error:
                st.caption("Los nombres y carnés no están disponibles con los permisos actuales; se muestran los identificadores.")
            f1, f2, f3 = st.columns(3)
            with f1:
                status_filter = st.selectbox("Estado", ["Todos", *states], key="admin_queue_status")
            with f2:
                careers = sorted({str(r.get("career_code")) for r in dashboard if r.get("career_code")})
                career_filter = st.selectbox("Carrera", ["Todas", *careers], key="admin_queue_career")
            with f3:
                search_filter = st.text_input("Buscar por nombre, carné o ID", key="admin_queue_search").strip().casefold()
            d1, d2 = st.columns(2)
            with d1:
                start_date = st.date_input("Desde (Guatemala)", value=None, key="admin_queue_from")
            with d2:
                end_date = st.date_input("Hasta (Guatemala)", value=None, key="admin_queue_to")
            visible = [r for r in dashboard
                if (status_filter == "Todos" or r["estado_panel"] == status_filter)
                and (career_filter == "Todas" or r.get("career_code") == career_filter)
                and (not search_filter or any(search_filter in str(field or "").casefold() for field in
                    (r["id"], r.get("student_id"), r["nombre_estudiante"], r["carne_estudiante"])))
                and (not start_date or (r["fecha_local"] and r["fecha_local"].date() >= start_date))
                and (not end_date or (r["fecha_local"] and r["fecha_local"].date() <= end_date))]
            # Priorizar pendientes antiguos, luego evaluación en curso y resultados cerrados.
            priority = {"Pendiente": 0, "En evaluación": 1, "Evaluado": 2, "Publicado": 3}
            visible.sort(key=lambda r: (priority.get(r["estado_panel"], 9),
                r["fecha_local"] or datetime.max.replace(tzinfo=guatemala_tz)))
            st.caption(f"{len(visible)} simulacros encontrados de {len(dashboard)} registrados (máximo 200 recientes).")
            pending_answers = sum(r["pendientes"] + r["borradores"] for r in dashboard)
            completed_count = sum(r["estado_panel"] == "Publicado" for r in dashboard)
            publish_pct = 100 * completed_count / len(dashboard) if dashboard else 0
            evaluated = [r for r in dashboard if r["estado_panel"] in ("Evaluado", "Publicado") and r["fecha_local"]]
            k1, k2, k3 = st.columns(3)
            k1.metric("Respuestas por completar", pending_answers)
            k2.metric("Simulacros publicados (%)", f"{publish_pct:.1f}%")
            k3.metric("Listos para publicar", sum(r["estado_panel"] == "Evaluado" for r in dashboard))
            if not visible:
                st.info("No hay simulacros que coincidan con los filtros.")
                return
            # El resultado solo se muestra si la publicación es oficial.
            published_ids = [r["id"] for r in visible if r["estado_panel"] == "Publicado"]
            official_scores = {}
            if published_ids:
                try:
                    for batch_start in range(0, len(published_ids), 100):
                        published = (grading_admin.client().table("attempts")
                            .select("id,score,grading_status,published_at,published_by")
                            .in_("id", published_ids[batch_start:batch_start+100])
                            .eq("grading_status", "publicado").execute().data or [])
                        official_scores.update({item["id"]: item for item in published})
                except Exception:
                    st.caption("No se pudieron recuperar las notas publicadas para el reporte.")
            table_rows = []
            for r in visible:
                official = official_scores.get(r["id"], {})
                table_rows.append({
                    "Simulacro": r["id"],
                    "Nombre": r["nombre_estudiante"],
                    "Carné": r["carne_estudiante"] or "No disponible",
                    "Estudiante (ID)": r.get("student_id") or "",
                    "Carrera": r.get("career_code") or "—",
                    "Finalizado (Guatemala)": r["fecha_local"].strftime("%d/%m/%Y %H:%M") if r["fecha_local"] else "—",
                    "Estado": r["estado_panel"],
                    "Evaluadas": f"{r['completas']}/{r['respuestas']}",
                    "Borradores": r["borradores"],
                    "Pendientes": r["pendientes"],
                    "Avance (%)": r["avance"],
                    "Nota oficial": float(official["score"]) if official.get("score") is not None else None,
                    "Publicado (UTC)": official.get("published_at") or "",
                    "Publicado por (ID)": official.get("published_by") or "",
                })
            report = pd.DataFrame(table_rows)
            st.dataframe(report.drop(columns=["Estudiante (ID)", "Publicado por (ID)"]).assign(
                Simulacro=lambda df: df["Simulacro"].str[:8]), use_container_width=True, hide_index=True)
            # Exportación nativa Excel con filtros, formato y hoja de resumen.
            from io import BytesIO
            from openpyxl import Workbook
            from openpyxl.styles import Font, PatternFill, Alignment
            from openpyxl.utils import get_column_letter
            from openpyxl.worksheet.table import Table, TableStyleInfo

            excel_buffer = BytesIO()
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "Evaluaciones"
            sheet.append(list(report.columns))
            for values in report.itertuples(index=False, name=None):
                sheet.append([None if pd.isna(v) else v for v in values])
            sheet.freeze_panes = "A2"
            sheet.sheet_view.showGridLines = False
            for cell in sheet[1]:
                cell.font = Font(name="Aptos", bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
                cell.alignment = Alignment(vertical="center", wrap_text=True)
            sheet.row_dimensions[1].height = 32
            if len(report):
                last_col = get_column_letter(len(report.columns))
                excel_table = Table(displayName="TablaEvaluaciones", ref=f"A1:{last_col}{len(report)+1}")
                excel_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2",
                    showFirstColumn=False, showLastColumn=False,
                    showRowStripes=True, showColumnStripes=False)
                sheet.add_table(excel_table)
            for col_idx, name in enumerate(report.columns, 1):
                width = min(44, max(14, len(str(name)) + 3))
                if name in ("Nombre", "Estudiante (ID)", "Simulacro"):
                    width = 38
                if name in ("Finalizado (Guatemala)", "Publicado (UTC)"):
                    width = 27
                sheet.column_dimensions[get_column_letter(col_idx)].width = width
            if "Nota oficial" in report.columns:
                score_col = report.columns.get_loc("Nota oficial") + 1
                for row_num in range(2, len(report) + 2):
                    sheet.cell(row_num, score_col).number_format = '0.00'
            summary = workbook.create_sheet("Resumen")
            summary.append(["PREP-EIME · Resumen de evaluaciones"])
            summary.merge_cells("A1:B1")
            summary["A1"].font = Font(name="Aptos Display", size=16, bold=True, color="17365D")
            summary.append(["Indicador", "Cantidad"])
            for state in ("Pendiente", "En evaluación", "Evaluado", "Publicado"):
                summary.append([state, sum(r["estado_panel"] == state for r in visible)])
            summary.append(["Simulacros incluidos", len(visible)])
            summary.append(["Respuestas pendientes", sum(r["pendientes"] + r["borradores"] for r in visible)])
            summary.column_dimensions["A"].width = 38
            summary.column_dimensions["B"].width = 20
            summary.freeze_panes = "A3"
            for cell in summary[2]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="17365D")
            workbook.save(excel_buffer)
            st.download_button("Descargar reporte en Excel (.xlsx)",
                data=excel_buffer.getvalue(),
                file_name="PREP_EIME_evaluaciones_administrativas.xlsx",
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=False)
            selected_id = st.selectbox("Abrir simulacro", [r["id"] for r in visible],
                format_func=lambda aid: next(
                    f"{r['id'][:8]} · {r['nombre_estudiante']} · {r.get('career_code') or '—'} · {r['estado_panel']} · {r['completas']}/{r['respuestas']}"
                    for r in visible if r["id"] == aid), key="admin_queue_selected")
            attempt = next(r for r in attempts if r["id"] == selected_id)
        except Exception as dashboard_error:
            st.warning(f"No se pudo cargar el panel: {dashboard_error}")
            idx=st.selectbox("Simulacro",range(len(attempts)),format_func=lambda i:
                f"{attempts[i]['id'][:8]} · {attempts[i].get('career_code','')} · {attempts[i].get('grading_status','pendiente')}")
            attempt=attempts[idx]
        tracking=grading_admin.progress(attempt["id"])
        st.subheader("Seguimiento de la evaluación")
        m1,m2,m3,m4=st.columns(4)
        m1.metric("Respuestas",tracking["total"])
        m2.metric("Completas",tracking["completas"])
        m3.metric("Borradores",tracking["borradores"])
        m4.metric("Pendientes",tracking["pendientes"])
        st.progress(tracking["porcentaje"]/100.0,
            text=f"Avance de evaluación: {tracking['porcentaje']:.1f}%")
        if tracking["total"] and tracking["completas"]==tracking["total"]:
            st.success("Todas las respuestas están evaluadas.")
            preview=grading_admin.publication_preview(attempt["id"])
            if preview["ready"]:
                if attempt.get("grading_status") == "publicado":
                    official_row = (grading_admin.client().table("attempts")
                        .select("score,published_at,published_by")
                        .eq("id", attempt["id"]).single().execute().data or {})
                    official_score = official_row.get("score")
                    st.metric("Nota oficial publicada",
                        f"{float(official_score):.2f}/100" if official_score is not None else "No disponible")
                    st.success("Este simulacro ya tiene resultados publicados.")
                    st.caption(f"Publicado: {official_row.get('published_at') or 'Sin fecha'}")
                else:
                    st.metric("Nota general prevista (solo administrador)",f"{preview['score']:.2f}/100")
                    st.caption("Promedio simple por respuesta; Supabase calcula la nota definitiva al publicar.")
                    # Solo permitir publicación cuando aún no está publicado.
                    ready_to_publish = grading_admin.publication_enabled()
                    if not ready_to_publish:
                        st.warning("Ejecuta la migración 131 y valida la autorización antes de activar ENABLE_ADMIN_PUBLICATION=true en Streamlit Secrets.")
                    confirmed = st.checkbox(
                        "Confirmo que revisé las evaluaciones y deseo publicar la nota oficial",
                        key=f"publish_confirm_{attempt['id']}",
                        disabled=not ready_to_publish,
                    )
                    if st.button(
                        "Publicar resultados oficialmente",
                        type="primary",
                        disabled=not (ready_to_publish and confirmed),
                        key=f"publish_{attempt['id']}",
                    ):
                        result = grading_admin.publish(attempt["id"])
                        st.success(f"Resultados publicados: {result}")
                        st.rerun()
            else:
                st.error("No se pudo verificar la integridad de las cinco puntuaciones por respuesta.")
        else:
            st.info("Complete las cinco valoraciones de cada respuesta antes de publicar.")
        answers=get_attempt_answers_v16(attempt["id"])
        if not answers:
            st.info("El simulacro no tiene respuestas")
            return
        def answer_label(i):
            row=answers[i]
            ev=tracking["por_respuesta"].get(row["id"],{})
            state={"completa":"Completa","borrador":"Borrador"}.get(ev.get("status"),"Pendiente")
            return f"{i+1}. {row.get('course_name','Curso')} · {state} · {row['id'][:8]}"
        # Mantener la respuesta seleccionada a través de st.rerun() tras guardar.
        # La selección se guarda por ID y por simulacro, no por posición global.
        selection_key = f"admin_selected_answer_{attempt['id']}"
        available_ids = [row["id"] for row in answers]
        if st.session_state.get(selection_key) not in available_ids:
            st.session_state[selection_key] = available_ids[0]
        selected_answer_id = st.selectbox(
            "Respuesta", available_ids,
            format_func=lambda answer_id: answer_label(available_ids.index(answer_id)),
            key=selection_key,
        )
        a = answers[available_ids.index(selected_answer_id)]
        st.markdown("**Pregunta**")
        st.write(a.get("question_text") or "Sin enunciado")
        st.markdown("**Respuesta del estudiante**")
        st.write(a.get("answer_text") or "Sin respuesta")
        with st.expander("Solución esperada"):
            st.write(a.get("expected_answer") or "No registrada")
        previous,scores=grading_admin.evaluation(a["id"])
        mapping={r["code"]:r for r in rubrics}
        prior=next((r["code"] for r in rubrics if previous and r["id"]==previous["rubric_id"]),rubrics[0]["code"])
        code=st.selectbox("Rúbrica",list(mapping),index=list(mapping).index(prior),
            format_func=lambda x:f"{x} · {mapping[x]['name']}",key=f"rubric_{a['id']}")
        criteria=grading_admin.criteria(mapping[code]["id"])
        saved={int(x["criterion_order"]):x["performance_level"] for x in scores} if code==prior else {}
        factors={"excelente":1.0,"bueno":0.8,"regular":0.6,"deficiente":0.3,"no_demostrado":0.0}
        items=[]
        total=0.0
        with st.form(f"admin_grade_{a['id']}_{code}"):
            st.caption(f"Estado: {(previous or {}).get('status','Nuevo')} · Versión: {(previous or {}).get('version','—')}")
            for criterion in criteria:
                n=int(criterion["criterion_order"])
                opts=["sin_evaluar"]+list(factors)
                val=saved.get(n,"sin_evaluar")
                level=st.selectbox(f"{n}. {criterion['name']} · máximo {criterion['max_points']} puntos",
                    opts,index=opts.index(val) if val in opts else 0,
                    format_func=lambda x:x.replace('_',' ').capitalize(),key=f"grade_{a['id']}_{code}_{n}")
                if level!="sin_evaluar":
                    items.append({"order":n,"level":level})
                    total+=float(criterion["max_points"])*factors[level]
            st.metric("Puntos parciales (no publicados)",f"{total:.2f}/100")
            st.caption(f"{len(items)} de 5 criterios. 'Sin evaluar' es distinto de 'No demostrado'.")
            submit=st.form_submit_button("Guardar evaluación",type="primary",
                disabled=attempt.get("grading_status")=="publicado")
        if submit:
            result=grading_admin.save(a["id"],code,items,previous["version"] if previous else None)
            st.success(f"Guardado correctamente: {result}")
            st.rerun()
    except Exception as exc:
        st.error(f"No se pudo completar la operación: {exc}")

def users_view(user):
    page_header("Usuarios y cohortes", "Administración de acceso")
    st.info(
        "La administración masiva de estudiantes será una etapa posterior, "
        "con importación desde Excel."
    )


def analytics_view(user):
    page_header("Analítica", "Indicadores iniciales")

    programs = get_programs(include_drafts=True)
    questions = get_all_questions(include_drafts=True)

    c1, c2, c3, c4 = st.columns(4)
    with c1:
        metric_card(
            "Cursos en catálogo",
            len(get_all_courses()),
            "Base real",
        )
    with c2:
        metric_card(
            "Programas registrados",
            len(programs),
            "Base real",
        )
    with c3:
        metric_card(
            "Preguntas EGP",
            len(questions),
            "Banco real",
        )
    with c4:
        metric_card(
            "Preguntas publicadas",
            sum(1 for q in questions if q["status"] == "publicado"),
            "Disponibles al estudiante",
        )


if "user" not in st.session_state:
    login_view()
    st.stop()

user = st.session_state.user
section = sidebar(user)

if section == "Inicio":
    home_view(user)
elif section == "Preparación académica":
    academic_view(user)
elif section == "Biblioteca":
    library_view(user)
elif section == "Preparación EGP":
    egp_view(user)
elif section == "Simulador EGP":
    if st.session_state.get("review_attempt_id"):
        attempt_review_view(user)
    else:
        simulator_view(user)
elif section == "Mi progreso":
    progress_view(user)
elif section == "Gestión académica":
    management_view(user)
elif section == "Banco de preguntas":
    question_bank_view(user)
elif section == "Control de calidad":
    quality_management_view(user)
elif section == "Revisión de respuestas":
    review_answers_view(user)
elif section == "Evaluación administrativa":
    administrative_grading_view(user)
elif section == "Usuarios y cohortes":
    users_view(user)
elif section == "Analítica":
    analytics_view(user)
