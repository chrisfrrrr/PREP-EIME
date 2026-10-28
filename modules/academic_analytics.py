"""Analítica administrativa de solo lectura sobre la bandeja existente.
No amplía permisos ni consulta puntuaciones de evaluaciones sin publicar.
"""
from io import BytesIO
import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter


def _export_student(rows, display_name):
    book = Workbook()
    ws = book.active
    ws.title = 'Expediente'
    headers = ['Simulacro', 'Carrera', 'Fecha Guatemala', 'Estado',
               'Respuestas completas', 'Respuestas totales', 'Nota oficial']
    ws.append(headers)
    for r in rows:
        ws.append([r['Simulacro'], r['Carrera'], r['Finalizado (Guatemala)'],
                   r['Estado'], r['Evaluadas'], r['Total respuestas'], r['Nota oficial']])
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    for cell in ws[1]:
        cell.fill = PatternFill('solid', fgColor='17365D')
        cell.font = Font(bold=True, color='FFFFFF')
        cell.alignment = Alignment(wrap_text=True)
    for i, width in enumerate([39, 25, 25, 21, 23, 20, 18], 1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cells in ws.iter_rows(min_row=2):
        cells[-1].number_format = '0.00'
    summary = book.create_sheet('Resumen')
    summary.append(['Expediente académico', display_name])
    summary.append(['Simulacros', len(rows)])
    published = [r['Nota oficial'] for r in rows if r['Nota oficial'] is not None]
    summary.append(['Publicados con nota', len(published)])
    summary.append(['Promedio de notas publicadas', sum(published)/len(published) if published else None])
    summary['B4'].number_format = '0.00'
    summary.column_dimensions['A'].width = 35
    summary.column_dimensions['B'].width = 43
    buffer = BytesIO()
    book.save(buffer)
    return buffer.getvalue()


def render_administrative_academics(dashboard, official_scores):
    """Muestra expediente y agregados de hasta 200 intentos administrativos recientes.

    official_scores solo incluye intentos publicados. No utiliza notas preliminares.
    """
    if not dashboard:
        st.info('Todavía no hay registros para analizar.')
        return
    rows = []
    for r in dashboard:
        official = official_scores.get(r['id'], {}) if r['estado_panel'] == 'Publicado' else {}
        score = official.get('score')
        rows.append({
            'Simulacro': r['id'],
            'Estudiante ID': r.get('student_id') or '',
            'Nombre': r.get('nombre_estudiante') or 'Sin nombre disponible',
            'Carné': r.get('carne_estudiante') or 'No disponible',
            'Carrera': r.get('career_code') or 'Sin carrera',
            'Finalizado (Guatemala)': (r['fecha_local'].strftime('%d/%m/%Y %H:%M')
                                      if r.get('fecha_local') else 'Sin fecha'),
            'Fecha orden': r.get('fecha_local'),
            'Estado': r['estado_panel'],
            'Evaluadas': int(r['completas']),
            'Total respuestas': int(r['respuestas']),
            'Nota oficial': float(score) if score is not None else None,
        })
    frame = pd.DataFrame(rows)
    published = frame[frame['Nota oficial'].notna()].copy()
    st.caption('Alcance: hasta 200 simulacros administrativos finalizados recientes. '
               'Las notas históricas de autoevaluación no forman parte de este análisis.')
    tab_students, tab_careers = st.tabs(['Expediente individual', 'Analítica académica'])
    with tab_students:
        students = (frame[['Estudiante ID','Nombre','Carné','Carrera']]
                    .drop_duplicates(subset=['Estudiante ID'])
                    .sort_values(['Nombre','Carné']))
        ids = students['Estudiante ID'].tolist()
        labels = {r['Estudiante ID']: f"{r['Nombre']} · {r['Carné']} · {r['Carrera']}"
                  for _, r in students.iterrows()}
        chosen = st.selectbox('Seleccionar estudiante', ids,
                              format_func=lambda sid: labels.get(sid, sid),
                              key='academic_student_id')
        student = frame[frame['Estudiante ID'] == chosen].copy()
        student = student.sort_values('Fecha orden', ascending=False, na_position='last')
        notes = student['Nota oficial'].dropna()
        m1, m2, m3, m4 = st.columns(4)
        m1.metric('Simulacros', len(student))
        m2.metric('Publicados', int(student['Nota oficial'].notna().sum()))
        m3.metric('Pendientes de publicación', int(student['Nota oficial'].isna().sum()))
        m4.metric('Promedio publicado', f'{notes.mean():.2f}/100' if len(notes) else 'Sin datos')
        st.dataframe(student[['Simulacro','Carrera','Finalizado (Guatemala)',
                              'Estado','Evaluadas','Total respuestas','Nota oficial']],
                     hide_index=True, use_container_width=True)
        if len(notes) >= 2:
            history = student[student['Nota oficial'].notna()].sort_values('Fecha orden')
            st.markdown('**Evolución de resultados publicados**')
            st.line_chart(history.set_index('Finalizado (Guatemala)')['Nota oficial'])
        st.download_button('Descargar expediente en Excel (.xlsx)',
            _export_student(student.to_dict('records'), labels[chosen]),
            file_name='PREP_EIME_expediente_estudiante.xlsx',
            mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
            key='download_academic_student')
    with tab_careers:
        c1, c2, c3 = st.columns(3)
        c1.metric('Estudiantes distintos', frame['Estudiante ID'].nunique())
        c2.metric('Simulacros administrativos', len(frame))
        c3.metric('Con nota publicada', len(published))
        summary = (frame.groupby('Carrera', dropna=False)
                   .agg(Simulacros=('Simulacro','count'),
                        Estudiantes=('Estudiante ID','nunique'),
                        Publicados=('Nota oficial','count'),
                        Promedio_publicado=('Nota oficial','mean')).reset_index())
        summary = summary.rename(columns={'Promedio_publicado':'Promedio publicado'})
        st.dataframe(summary, use_container_width=True, hide_index=True)
        if not published.empty:
            st.markdown('**Promedio de notas publicadas por carrera**')
            means = published.groupby('Carrera')['Nota oficial'].mean().sort_values()
            st.bar_chart(means)
            st.caption('Los promedios excluyen simulacros pendientes y sin nota. '
                       'No representan el rendimiento de todos los estudiantes.')
        else:
            st.info('Todavía no hay calificaciones administrativas publicadas.')
